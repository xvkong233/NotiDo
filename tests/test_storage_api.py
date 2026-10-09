import io
from types import SimpleNamespace

import pytest

from notido.api import PagesAPI
from notido.db import Database, execute
from notido.errors import NotiDoError
from notido.materials import AsyncFile, BlobStore
from notido.sources import DisabledWebsiteSource, SourceRegistry


async def test_late_original_requires_exact_saved_task_and_deduplicates(service):
    import hashlib

    from conftest import create_plan, operation, parsed

    service.bridge.plan = create_plan
    group = await parsed(service)
    first = await operation(service)
    await service.execute_operation({"operation_id": first["id"]})
    api = PagesAPI(service)
    payload = b"late original\x00bytes"

    def request(revision, request_id, target):
        async def form():
            return {}

        async def files():
            return {"file": AsyncFile(io.BytesIO(payload))}

        return SimpleNamespace(
            username="admin",
            form=form,
            files=files,
            query={
                "request_id": request_id,
                "expected_revision": str(revision),
                "sha256": hashlib.sha256(payload).hexdigest(),
                "name": "late.zip",
                "target_operation_id": target,
            },
        )

    current = (await service.db.read("SELECT * FROM material_groups"))[0]
    with pytest.raises(NotiDoError) as error:
        await api.upload_file(request(current["revision"], "unspecified", None), group["group_id"])
    assert error.value.code == "LATE_MATERIAL_TARGET_REQUIRED"
    value = await api.upload_file(
        request(current["revision"], "exact", first["id"]), group["group_id"]
    )
    assert value["attachment_scheduled"]
    assert value == await api.upload_file(
        request(current["revision"], "exact", first["id"]), group["group_id"]
    )
    # A new occurrence of the same bytes still targets the one hash operation.
    await api.upload_file(
        request(value["revision"], "new-occurrence", first["id"]), group["group_id"]
    )
    assert len(await service.db.read("SELECT * FROM assets")) == 2
    assert len(await service.db.read("SELECT * FROM operations WHERE kind='upload'")) == 1
    assert len(service.gateway.items) == 1


async def test_blob_bytes_immutable_and_duplicate_source_retained(tmp_path):
    store = BlobStore(tmp_path)
    first = await store.save(AsyncFile(io.BytesIO(b"original\x00bytes")), min_free_bytes=0)
    second = await store.save(AsyncFile(io.BytesIO(b"original\x00bytes")), min_free_bytes=0)
    assert first == second and store.path(first["path"]).read_bytes() == b"original\x00bytes"
    different = await store.save(AsyncFile(io.BytesIO(b"different")), min_free_bytes=0)
    assert different["path"] != first["path"]
    assert not list((tmp_path / "staging").iterdir())


async def test_blob_limits_and_path_traversal(tmp_path):
    store = BlobStore(tmp_path)
    with pytest.raises(NotiDoError):
        await store.save(AsyncFile(io.BytesIO(b"too long")), max_bytes=2, min_free_bytes=0)
    assert not list((tmp_path / "blobs").iterdir())
    with pytest.raises(NotiDoError):
        store.path("../arbitrary")


async def test_pdf_docx_and_txt_read_real_bytes(tmp_path):
    from docx import Document
    from pypdf import PdfWriter

    store = BlobStore(tmp_path)
    text = tmp_path / "text.txt"
    text.write_text("2027年12月20日提交 PDF。", encoding="utf-8")
    result = await store.read(text, "text.txt")
    assert "2027" in result["segments"][0]["text"] and not result["unknowns"]
    doc = Document()
    doc.add_paragraph("第一段")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "第二项", "提交格式 PDF"
    doc.add_paragraph("第三段")
    target = tmp_path / "sample.docx"
    doc.save(target)
    result = await store.read(target, "sample.docx")
    assert [x["text"] for x in result["segments"]] == ["第一段", "第二项\t提交格式 PDF", "第三段"]
    pdf = PdfWriter()
    for _ in range(31):
        pdf.add_blank_page(width=100, height=100)
    pdf.write(tmp_path / "long.pdf")
    result = await store.read(tmp_path / "long.pdf", "long.pdf")
    assert any(x.startswith("PDF_PAGE_LIMIT") for x in result["unknowns"])


async def test_unsupported_original_not_unzipped(tmp_path):
    store = BlobStore(tmp_path)
    original = tmp_path / "dangerous.zip"
    original.write_bytes(b"not even a zip; still an original attachment")
    result = await store.read(original, "dangerous.zip")
    assert result["unknowns"] == ["BODY_NOT_SUPPORTED_ATTACHMENT_ONLY"]
    assert not result["segments"]


async def test_api_dedup_precedes_revision_and_rejects_reused_payload(service):
    calls = []

    async def mutation(conn, revision):
        calls.append(revision)
        await execute(conn, "UPDATE settings SET revision=revision+1")
        return {"revision": revision + 1}

    result = await service.db.api_mutate("test", "main", "r", 0, {"v": 1}, mutation)
    assert result == await service.db.api_mutate("test", "main", "r", 0, {"v": 1}, mutation)
    assert calls == [0]
    with pytest.raises(NotiDoError) as error:
        await service.db.api_mutate("test", "main", "r", 1, {"v": 2}, mutation)
    assert error.value.code == "REQUEST_ID_REUSED"


def test_web_principal_required_on_every_wrapper():
    for request in (
        SimpleNamespace(username=None, plugin_name="astrbot_plugin_notido"),
        SimpleNamespace(username="admin", plugin_name="another"),
    ):
        with pytest.raises(NotiDoError):
            PagesAPI.authenticate(request)
    PagesAPI.authenticate(SimpleNamespace(username="admin", plugin_name="astrbot_plugin_notido"))


async def test_auth_commit_failure_sets_persistent_maintenance(service, monkeypatch):
    api = PagesAPI(service)
    service.gateway.runner = SimpleNamespace(home=service.root / "cli-home")
    config = service.gateway.runner.home / ".config/dida-cli/config.json"
    config.parent.mkdir(parents=True)
    config.write_text('{"access_token":"local-test-only"}')

    async def json_body():
        return {"request_id": "auth-clear", "expected_revision": 0, "confirm": "清除当前授权"}

    async def failed_commit(endpoint, target, body, mutation):
        async with service.db.transaction() as conn:
            await mutation(conn, 0)
            raise RuntimeError("simulated DB failure after file switch")

    monkeypatch.setattr(api, "mutate", failed_commit)
    with pytest.raises(RuntimeError):
        await api.clear_auth(SimpleNamespace(json=json_body))
    assert service.maintenance == "AUTH_COMMIT_UNKNOWN"
    assert (service.root / "maintenance.required").read_text().strip() == "AUTH_COMMIT_UNKNOWN"
    assert (await service.db.settings())[0].account_ref == "account-a"
    assert not config.exists()


async def test_schema_too_new_and_checksum_fail(tmp_path):
    db = Database(tmp_path)
    await db.initialize()
    async with db.transaction() as conn:
        await execute(conn, "UPDATE schema_migrations SET checksum='modified'")
    with pytest.raises(NotiDoError) as error:
        await db.initialize()
    assert error.value.code == "MIGRATION_CHANGED"
    async with db.transaction() as conn:
        await execute(conn, "INSERT INTO schema_migrations VALUES (999,'future',0)")
    with pytest.raises(NotiDoError) as error:
        await db.initialize()
    assert error.value.code == "SCHEMA_TOO_NEW"
    await db.close()


async def test_foreign_keys_enforced_and_account_scope_unique(service):
    from sqlalchemy.exc import IntegrityError

    with pytest.raises(IntegrityError):
        async with service.db.transaction() as conn:
            await execute(
                conn, "INSERT INTO account_scopes VALUES ('bad','missing','cn',NULL,0,'active',0)"
            )
    with pytest.raises(IntegrityError):
        async with service.db.transaction() as conn:
            await execute(
                conn, "INSERT INTO account_scopes VALUES ('b','personal','cn',NULL,0,'active',0)"
            )


async def test_disabled_sources_never_network():
    registry = SourceRegistry()
    assert set(registry.adapters) == {"website"}
    assert registry.describe()[0]["supported"] is False
    with pytest.raises(NotiDoError) as error:
        registry.enable("website")
    assert error.value.code == "SOURCE_NOT_IMPLEMENTED"
    with pytest.raises(NotiDoError):
        await DisabledWebsiteSource().poll()


async def test_memory_source_delivers_same_uniform_envelope_without_production_registration(
    service,
):
    from conftest import envelope

    from notido.sources import MemorySourceAdapter, SourceEnvelope

    sample = SourceEnvelope(
        envelope("memory-sample", text="脱敏内存通知", source_kind="user_forward"),
        "test-only:sample",
    )
    adapter = MemorySourceAdapter([sample])
    values = await adapter.poll()
    result = await service.intake(values[0].envelope)
    assert result["group_id"] and values[0].source_ref == "test-only:sample"
    assert not await adapter.poll() and not await adapter.fetch_materials(sample.source_ref)
    assert "memory" not in service.sources.adapters


async def test_cancel_claim_condition_preserves_result(service):
    from conftest import create_plan, operation, parsed

    service.bridge.plan = create_plan
    await parsed(service)
    op = await operation(service)
    async with service.db.transaction() as conn:
        canceled = await execute(
            conn,
            "UPDATE operations SET state='cancelled' WHERE id=:id AND state='validated'",
            {"id": op["id"]},
        )
    assert canceled.rowcount == 1
    await service.execute_operation({"operation_id": op["id"]})
    assert not service.gateway.writes
