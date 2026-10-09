"""Opt-in real Provider material fixtures; each admitted message has a durable journal."""

import hashlib
import json
import subprocess
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

from tools.container_smoke import BASE, PREFIX, ROOT, request, unwrap, wait_ready
from tools.live_same_name_smoke import save


def fixture(kind, path, run_label=""):
    image = Image.new("RGB", (1800, 800), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 48)
    labels = {
        "png": ("截图", "24", "17:40", "ND-V01"),
        "pdf": ("扫描PDF", "25", "14:10", "ND-V02"),
        "docx": ("DOCX嵌图", "26", "09:25", "ND-V03"),
    }
    label, day, hour, code = labels[kind]
    code += run_label
    lines = [
        f"NotiDo 验收 · {label}通知",
        f"本人必做事项：提交 {code} 材料验收报告",
        f"截止时间：2027年12月{day}日{hour}",
        f"提交要求：PDF 格式，并填写验收编号 {code}",
        "请把这份通知原件附到报告任务。",
    ]
    for index, line in enumerate(lines):
        draw.text((70, 60 + index * 115), line, font=font, fill="black")
    if kind == "png":
        image.save(path)
    elif kind == "pdf":
        image.save(path, "PDF", resolution=150)
    else:
        original = path.with_suffix(".png")
        image.save(original)
        document = Document()
        document.add_paragraph(
            "这是我的必做通知，截止时间和提交要求都在下方嵌图中。请阅读嵌图并记录事项。"
        )
        document.add_picture(str(original), width=Inches(6.4))
        document.save(path)
    return {"date": f"2027-12-{day}", "time": hour, "code": code}


def upload(token, group, revision, path, request_id):
    data = path.read_bytes()
    boundary = "NotiDo" + uuid.uuid4().hex
    multipart = (
        (
            f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="fixture"\r\nContent-Type: application/octet-stream\r\n\r\n'
        ).encode()
        + data
        + f"\r\n--{boundary}--\r\n".encode()
    )
    values = {
        "request_id": request_id,
        "expected_revision": str(revision),
        "sha256": hashlib.sha256(data).hexdigest(),
        "name": path.name,
    }
    req = urllib.request.Request(
        BASE + PREFIX + f"groups/{group}/files?" + urllib.parse.urlencode(values),
        data=multipart,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            return unwrap(json.loads(response.read()))
    except urllib.error.HTTPError as error:
        value = unwrap(json.loads(error.read()))
        raise RuntimeError(
            f"fixture upload rejected: {value.get('error', {}).get('code', error.code)}"
        ) from None


def details(token, group):
    status, response = request(PREFIX + f"notices/{group}", token=token)
    if status != 200:
        raise RuntimeError("fixture group unavailable")
    detail = unwrap(response)
    status, response = request(PREFIX + "operations?limit=100", token=token)
    if status != 200:
        raise RuntimeError("operation list unavailable")
    operations = []
    for row in unwrap(response)["items"]:
        status, record = request(PREFIX + f"operations/{row['id']}", token=token)
        value = unwrap(record)
        if status == 200 and value["plan"].get("group_id") == group:
            operations.append(value)
    detail["operations"] = operations
    return detail


def main():
    import sys

    kind = sys.argv[1]
    if kind not in ("png", "pdf", "docx"):
        raise ValueError("png/pdf/docx required")
    token = wait_ready()
    private = ROOT / "runtime-data"
    run_label = sys.argv[2] if len(sys.argv) > 2 else ""
    if run_label and (not run_label.isalnum() or len(run_label) > 20):
        raise ValueError("short alphanumeric fixture run label required")
    suffix = f"-{run_label}" if run_label else ""
    path = private / f"NotiDo-{kind}-视觉原件{suffix}.{kind}"
    journal = private / f"container-vision-{kind}{suffix}-results.json"
    if journal.exists():
        state = json.loads(journal.read_text(encoding="utf-8"))
    else:
        expected = fixture(kind, path, run_label)
        state = {
            "message_id": "notido-vision-" + str(uuid.uuid4()),
            "expected": expected,
            "text": f"请记录我的必做事项，通知内容请读取随附{kind}原件，并把通知原件附到对应任务。",
            "upload_request": str(uuid.uuid4()),
            "close_request": str(uuid.uuid4()),
            "admitted_at": time.time(),
        }
        save(journal, state)
    if not state.get("group_id"):
        result = subprocess.run(
            [
                "docker",
                "exec",
                "-i",
                "-w",
                "/opt/notido-plugin",
                "-e",
                "PYTHONPATH=/AstrBot:/opt/notido-plugin",
                "notido-acceptance",
                "python",
                "-m",
                "tools.container_admit_fixture",
            ],
            input=json.dumps(state).encode(),
            capture_output=True,
            timeout=30,
        )
        if result.returncode:
            raise RuntimeError("public fixture admission failed")
        state["group_id"] = json.loads(result.stdout)["group_id"]
        save(journal, state)
    if not state.get("asset_id"):
        detail = details(token, state["group_id"])
        state.setdefault("upload_revision", detail["group"]["revision"])
        save(journal, state)
        value = upload(
            token, state["group_id"], state["upload_revision"], path, state["upload_request"]
        )
        state["asset_id"] = value["asset_id"]
        save(journal, state)
    if not state.get("closed"):
        detail = details(token, state["group_id"])
        if detail["group"]["state"] == "collecting":
            state.setdefault("close_revision", detail["group"]["revision"])
            save(journal, state)
            status, value = request(
                PREFIX + f"groups/{state['group_id']}/close",
                token=token,
                payload={
                    "request_id": state["close_request"],
                    "expected_revision": state["close_revision"],
                },
            )
            if status != 200 or "error" in unwrap(value):
                raise RuntimeError("fixture close failed")
        state["closed"] = True
        state["closed_at"] = time.time()
        save(journal, state)
    deadline = time.monotonic() + 65
    while time.monotonic() < deadline:
        detail = details(token, state["group_id"])
        operations = detail["operations"]
        creates = [x for x in operations if x["kind"] == "create"]
        uploads = [x for x in operations if x["kind"] == "upload"]
        if detail["group"]["state"] == "awaiting_clarification" or any(
            x["state"]
            not in (
                "validated",
                "executing",
                "created_unverified",
                "uploaded_unverified",
                "applied_unverified",
                "succeeded",
            )
            for x in operations
        ):
            state["detail"] = detail
            save(journal, state)
            raise RuntimeError(f"real {kind} fixture paused; inspect private evidence, no replay")
        if (
            len(creates) == 1
            and len(uploads) == 1
            and all(x["state"] == "succeeded" for x in operations)
        ):
            task, attachment = creates[0], uploads[0]
            expected, date = state["expected"], task["plan"]["normalized_date"]
            actual = json.loads(task["result"])["actual_fields"]
            if (
                date["local_date"] != expected["date"]
                or date["local_time"] != expected["time"]
                or expected["code"] not in actual.get("content", "")
                or "PDF" not in actual.get("content", "")
            ):
                raise RuntimeError("real visual deadline or requirements mismatch")
            if (
                attachment["plan"]["task_id"] != task["remote_id"]
                or attachment["plan"]["asset_id"] != state["asset_id"]
                or any(x["attempt"] != 1 for x in operations)
            ):
                raise RuntimeError("native original relation or attempts mismatch")
            state.update(
                {
                    "detail": detail,
                    "completed_at": time.time(),
                    "resumed": bool(state.get("resume_requests")),
                    "checks": [
                        {
                            "check": f"real_{kind}_visual_evidence_deadline_requirements",
                            "pass": True,
                        },
                        {"check": f"real_{kind}_original_native_hash_target", "pass": True},
                        {"check": f"real_{kind}_single_create_and_upload", "pass": True},
                    ],
                }
            )
            save(journal, state)
            print(json.dumps(state["checks"]))
            return
        time.sleep(0.5)
    raise RuntimeError("same fixture still pending; rerun only observes it")


if __name__ == "__main__":
    main()
