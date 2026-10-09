"""Journaled native AI evidence for late pages, incomplete and related materials.

Synthetic files contain no credentials. All model and task operations use the
authorized local AstrBot and dedicated Dida list. Interrupted requests are never
resent; failures and synthetic remote tasks remain for inspection.
"""

import argparse
import hashlib
import json
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from tools import native_corpus_smoke as collector
from tools.container_smoke import PREFIX, ROOT, request, unwrap, wait_ready
from tools.native_material_smoke import upload
from tools.native_performance_smoke import bind_session

JOURNAL = ROOT / "runtime-data/native-complex-material-v1-results.json"
FIXTURES = ROOT / "runtime-data/native-complex-material-v1"
BATCH = "v1"


def independent_text(text):
    for name in (
        "后页验收报告",
        "表格嵌图报告",
        "文本验收报告",
        "模糊时间报告",
        "个人实践报告",
        "关系报告甲",
        "关系报告乙",
        "超页报告",
    ):
        text = text.replace(name, "独立" + BATCH.upper() + name)
    if BATCH in ("v4", "v5"):
        text = text.replace("领取独立材料", "领取独立" + BATCH.upper() + "材料")
    return text


def picture(lines):
    if BATCH != "v1":
        lines = [*lines, "独立材料验收批次：" + BATCH.upper()]
    if BATCH in ("v3", "v4", "v5"):
        lines = [
            independent_text(line)
            .replace("ND-M01-LATE", "ND-M01-LATE-" + BATCH.upper())
            .replace("ND-M02-IMAGE", "ND-M02-IMAGE-" + BATCH.upper())
            for line in lines
        ]
    value = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(value)
    font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 42)
    for index, text in enumerate(lines):
        draw.text((60, 60 + index * 100), text, fill="black", font=font)
    return value


def fixtures():
    FIXTURES.mkdir(exist_ok=True)
    cases = []

    def case(key, files, expected, prompt="", unknown=False):
        cases.append(
            {
                "id": key,
                "files": [str(p) for p in files],
                "expected": expected,
                "prompt": prompt,
                "expected_unknown": unknown,
            }
        )

    path = FIXTURES / "M01-后页扫描通知.pdf"
    pages = [
        picture(
            [
                f"必修材料通知，第 {i} 页，共 6 页。",
                "本人必做：提交后页验收报告。",
                "具体截止时间与格式仅在第 6 页。",
                "前页背景不得代替第 6 页要求。",
            ]
        )
        for i in range(1, 6)
    ]
    pages.append(
        picture(
            [
                "必修材料通知，第 6 页，共 6 页。",
                "本人必做：提交后页验收报告。",
                "截止时间：2027年12月22日16:35",
                "提交格式：PDF；编号 ND-M01-LATE。",
                "通知原件附到这一个报告任务。",
            ]
        )
    )
    pages[0].save(path, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    case(
        "M01",
        [path],
        [
            {
                "title": "后页验收报告",
                "date": "2027-12-22",
                "time": "16:35",
                "requirements": ["PDF", "ND-M01-LATE"],
                "attachments": [path.name],
            }
        ],
    )

    path = FIXTURES / "M02-表格与嵌图.docx"
    image_path = FIXTURES / "M02-关键嵌图.png"
    picture(
        [
            "唯一截止时间：2027年12月23日09:45",
            "必填编号：ND-M02-IMAGE。",
            "段落、表格和本图均属于同一通知。",
        ]
    ).save(image_path)
    document = Document()
    document.add_paragraph("本人必须提交表格嵌图报告；原通知日期与编号只能从关键嵌图读取。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "提交渠道"
    table.rows[0].cells[1].text = "课程系统；文件格式PDF；命名为验收编号_报告.pdf"
    document.add_picture(str(image_path), width=Inches(6.4))
    document.save(path)
    case(
        "M02",
        [path],
        [
            {
                "title": "表格嵌图报告",
                "date": "2027-12-23",
                "time": "09:45",
                "requirements": ["PDF", "课程系统", "ND-M02-IMAGE"],
                "attachments": [path.name],
            }
        ],
    )

    path = FIXTURES / "M03-UTF8中文.txt"
    path.write_text(
        "本人必做：提交文本验收报告。截止2027年12月24日18:20。\n"
        "要求PDF格式，编号ND-M03-UTF8；请附此通知原件。",
        encoding="utf-8",
    )
    case(
        "M03",
        [path],
        [
            {
                "title": "文本验收报告",
                "date": "2027-12-24",
                "time": "18:20",
                "requirements": ["PDF", "ND-M03-UTF8"],
                "attachments": [path.name],
            }
        ],
    )

    path = FIXTURES / "M04-未知编码.txt"
    path.write_bytes(
        "本人必做：提交编码报告。截止2027年12月25日18:20。编号ND-M04。".encode("gb18030")
    )
    case("M04", [path], [], unknown=True)

    path = FIXTURES / "M05-损坏关键附件.pdf"
    path.write_bytes(b"%PDF-1.7\ntruncated-not-a-document\n")
    case(
        "M05",
        [path],
        [
            {
                "title": "领取独立材料",
                "date": "2027-12-26",
                "time": "10:15",
                "requirements": [],
                "attachments": [],
            }
        ],
        "独立明确事项：本人2027年12月26日10:15必须领取独立材料。"
        "另需提交损坏附件报告，其时间、要求只在随附文件，正文没有。两项互不依赖。",
        True,
    )

    path = FIXTURES / "M06-模糊关键时间.png"
    visual = picture(
        [
            "本人必做：提交模糊时间报告。",
            "截止时间：2027年12月27日16:25",
            "请以清晰原件中的时间为准，不猜测模糊数字。",
            "要求PDF格式，编号ND-M06。",
        ]
    )
    visual.paste(
        visual.crop((270, 145, 1600, 225)).filter(ImageFilter.GaussianBlur(40)), (270, 145)
    )
    visual.save(path)
    case("M06", [path], [])

    path = FIXTURES / "M07-混合身份.txt"
    path.write_text(
        "东北大学材料学院2024级所有学生必须提交个人实践报告，截止2027年12月28日15:10，PDF格式。\n"
        "只有班长还须汇总班级统计表，截止2027年12月29日11:20，Excel格式。"
        "两项相互独立；只给个人实践报告附本通知，班长身份未确认时询问受影响项。",
        encoding="utf-8",
    )
    case(
        "M07",
        [path],
        [
            {
                "title": "个人实践报告",
                "date": "2027-12-28",
                "time": "15:10",
                "requirements": ["PDF"],
                "attachments": [path.name],
            }
        ],
    )

    shared = FIXTURES / "M08-共用通知.txt"
    first, second = FIXTURES / "M08-专属甲.txt", FIXTURES / "M08-专属乙.txt"
    shared.write_text(
        "本人两项必做：提交关系报告甲，截止2027年12月29日14:20；"
        "提交关系报告乙，截止2027年12月30日09:15。均为PDF。\n"
        "本共用通知须分别附到两项任务。M08-专属甲.txt只属于报告甲；"
        "M08-专属乙.txt只属于报告乙。不得混挂。",
        encoding="utf-8",
    )
    first.write_text("关系报告甲的专属模板说明，必填编号ND-M08-A。", encoding="utf-8")
    second.write_text("关系报告乙的专属模板说明，必填编号ND-M08-B。", encoding="utf-8")
    case(
        "M08",
        [shared, first, second],
        [
            {
                "title": "关系报告甲",
                "date": "2027-12-29",
                "time": "14:20",
                "requirements": ["PDF", "ND-M08-A"],
                "attachments": [shared.name, first.name],
            },
            {
                "title": "关系报告乙",
                "date": "2027-12-30",
                "time": "09:15",
                "requirements": ["PDF", "ND-M08-B"],
                "attachments": [shared.name, second.name],
            },
        ],
    )

    path = FIXTURES / "M09-超页关键材料.pdf"
    first_page = picture(
        [
            "本人必做：提交超页报告。",
            "全文31页，唯一截止时间与格式位于第31页。",
            "当前页没有截止时间，不得按无日期保存。",
        ]
    )
    pages = [
        picture([f"超页报告背景，第{i}页，共31页。", "本页没有关键日期和要求。"])
        for i in range(2, 31)
    ]
    pages.append(picture(["第31页。截止2027年12月31日20:10，PDF格式，编号ND-M09。"]))
    first_page.save(path, "PDF", save_all=True, append_images=pages, resolution=120)
    case("M09", [path], [], unknown=True)
    if BATCH != "v1":
        from pathlib import Path

        for c in cases:
            for value in c["files"]:
                item = Path(value)
                if item.suffix == ".txt" and c["id"] != "M04":
                    with item.open("a", encoding="utf-8") as output:
                        output.write("\n独立材料验收批次：" + BATCH.upper())
                    if BATCH in ("v3", "v4", "v5"):
                        item.write_text(
                            independent_text(item.read_text(encoding="utf-8")), encoding="utf-8"
                        )
                elif item.suffix == ".docx":
                    doc = Document(item)
                    doc.add_paragraph("独立材料验收批次：" + BATCH.upper())
                    if BATCH in ("v3", "v4", "v5"):
                        doc.paragraphs[0].text = independent_text(doc.paragraphs[0].text)
                    doc.save(item)
        if BATCH in ("v4", "v5"):
            for c in cases:
                c["prompt"] = independent_text(c["prompt"])
                if c["id"] == "M05":
                    c["expected"][0]["title"] = "领取独立" + BATCH.upper() + "材料"
    return cases


def check(record):
    slot = record["evaluation"]
    rows = slot["operations"]
    # A native AI may update the unique existing synthetic task. Count the
    # current real task results, not only new rows, and never infer from args.
    creates = list({r["remote_id"]: r for r in rows if r["kind"] in ("create", "update")}.values())
    checks = {
        "exact_action_count": len(creates) == len(record["expected"]),
        "single_attempt_verified": all(
            r["state"] == "succeeded" and r["attempt"] == 1 for r in rows
        ),
        "no_other_mutations": all(r["kind"] in ("create", "update", "upload") for r in rows),
    }
    unmatched = creates.copy()
    for wanted in record["expected"]:
        row = next(
            (
                r
                for r in unmatched
                if wanted["title"]
                in json.loads(r["result"] or "{}").get("actual_fields", {}).get("title", "")
            ),
            None,
        )
        key = wanted["title"]
        checks[key + ":exists"] = row is not None
        if row is None:
            continue
        unmatched.remove(row)
        fields = json.loads(row["result"] or "{}").get("actual_fields", {})
        local = (
            datetime.fromisoformat(fields["dueDate"].replace("Z", "+00:00")).astimezone(
                ZoneInfo("Asia/Shanghai")
            )
            if fields.get("dueDate")
            else None
        )
        checks[key + ":deadline"] = bool(
            local
            and local.strftime("%Y-%m-%d %H:%M") == wanted["date"] + " " + wanted["time"]
            and fields.get("isAllDay") is False
        )
        checks[key + ":requirements"] = all(
            v in fields.get("content", "") for v in wanted["requirements"]
        )
        uploads = [
            r
            for r in rows
            if r["kind"] == "upload" and r["plan"].get("task_id") == row["remote_id"]
        ]
        hashes = [
            json.loads(r["result"] or "{}").get("actual_fields", {}).get("sha256") for r in uploads
        ]
        expected_hashes = [record["file_hashes"][name] for name in wanted["attachments"]]
        checks[key + ":attachment_hashes_and_targets"] = sorted(hashes) == sorted(expected_hashes)
    checks["no_orphan_uploads"] = all(
        r["plan"].get("task_id") in {c["remote_id"] for c in creates}
        for r in rows
        if r["kind"] == "upload"
    )
    return checks


def main(batch="v1"):
    global JOURNAL, FIXTURES, BATCH
    BATCH = batch
    JOURNAL = ROOT / f"runtime-data/native-complex-material-{batch}-results.json"
    FIXTURES = ROOT / f"runtime-data/native-complex-material-{batch}"
    collector.JOURNAL = JOURNAL
    token = wait_ready()
    settings = unwrap(request(PREFIX + "settings", token=token)[1])["settings"]
    project = json.loads((ROOT / "runtime-data/test-project.json").read_text(encoding="utf-8-sig"))
    if project["name"] != "NotiDo 验收" or settings["allowed_projects"] != [project["id"]]:
        raise RuntimeError("dedicated test scope required")
    installed = collector.runtime_version()
    if JOURNAL.exists():
        state = json.loads(JOURNAL.read_text(encoding="utf-8"))
    else:
        state = {
            **installed,
            "date": datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat(),
            "batch": batch,
            "cases": {c["id"]: {**c, "phase": "ready"} for c in fixtures()},
        }
        collector.save(state)
    if any(state.get(k) != v for k, v in installed.items()):
        raise RuntimeError("retain source version per acceptance batch")
    from pathlib import Path

    for record in state["cases"].values():
        if record["phase"] == "collected":
            continue
        if record["phase"] == "ready":
            record["session"] = bind_session(token)
            record["phase"] = "bound"
            collector.save(state)
        if record["phase"] == "bound":
            seeded = collector.send_once(
                token,
                state,
                record,
                "memory",
                "我是东北大学材料学院2024级学生，班长等其他角色未确认。"
                "唯一允许清单为NotiDo 验收。稍后直接处理本人必做通知；各任务标题加“NotiDo 验收 · 复杂材料 · "
                + record["id"]
                + (" · " + batch.upper() if batch != "v1" else "")
                + " · ”前缀并保留事项名称。不确定关键日期或身份，只追问受影响项；"
                "独立明确项直接记录。现在只记住这些会话条件，不调用变更工具。",
            )
            if seeded["operations"]:
                raise RuntimeError("memory must not mutate tasks")
            record["phase"] = "seeded"
            collector.save(state)
        if record["phase"] == "seeded":
            record.setdefault("attachments", {})
            record.setdefault("file_hashes", {})
            for value in record["files"]:
                path = Path(value)
                if path.name not in record["attachments"]:
                    record["attachments"][path.name] = upload(token, path)
                    record["file_hashes"][path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
                    collector.save(state)
            record["phase"] = "prepared"
            collector.save(state)
        if record["phase"] == "prepared":
            message = [
                {
                    "type": "plain",
                    "text": "这是本人必做通知，请使用材料工具读取全部可读位置，保存明确事项并把通知原件按所述关系挂到原生附件。"
                    "关键材料缺失或不可读时不要猜测。"
                    + ("本轮独立材料验收批次：" + batch.upper() + "。" if batch != "v1" else "")
                    + record["prompt"],
                }
            ]
            if batch in ("v3", "v4", "v5"):
                message[0]["text"] += (
                    "本批是另一组独立验收事项，与历史批次任务相互独立。只处理本轮前缀任务，不修改历史任务；通知中独立"
                    + batch.upper()
                    + "是事项标识，不是同事项的新版本。"
                )
            message.extend(
                {
                    "type": "image" if name.endswith(".png") else "file",
                    "attachment_id": a["attachment_id"],
                    "filename": a["filename"],
                }
                for name, a in record["attachments"].items()
            )
            collector.send_once(token, state, record, "evaluation", message)
            record["checks"] = check(record)
            record["phase"] = "collected"
            collector.save(state)
            print(
                json.dumps(
                    {
                        "case": record["id"],
                        "fields_pass": all(record["checks"].values()),
                        "semantic_review_pending": True,
                    }
                ),
                flush=True,
            )
    print(
        json.dumps(
            {
                "collected": len(state["cases"]),
                "fields_pass": sum(all(r["checks"].values()) for r in state["cases"].values()),
                "semantic_review_pending": True,
            }
        ),
        flush=True,
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--batch", choices=("v1", "v2", "v3", "v4", "v5"), default="v1")
    main(parser.parse_args().batch)
